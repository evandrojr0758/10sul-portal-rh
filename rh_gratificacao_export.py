"""Compartilhamento e exportação do retrato enviado para aprovação."""
from collections import defaultdict
from io import BytesIO
from math import isfinite, ceil
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


COLUNAS = ["COLABORADOR", "EMPRESA", "FUNÇÃO", "FRENTE", "SALÁRIO",
           "EQUIPE REVISÃO", "MÉDIA TEMPO (h)", "FALTAS", "ATESTADOS (dias)",
           "DNA", "DESVIOS", "INTEGRAL", "GRATIFICAÇÃO", "STATUS", "MOTIVO / CÁLCULO"]
MOEDA = '"R$" #,##0.00'
STATUS = {"AGUARDANDO_AUTORIZACAO": "AGUARDANDO AUTORIZAÇÃO", "AUTORIZADA": "AUTORIZADA"}


def valor_numero(v):
    try:
        f = float(v)
        return f if isfinite(f) else None
    except (ValueError, TypeError):
        return None


def whatsapp_aprovacao(base_url, token, ano, mes):
    link = base_url.split("?")[0].rstrip("/") + "/?aprovar_gratificacao=" + quote(str(token), safe="")
    mensagem = f"Olá! Segue a apuração de gratificação de {int(mes):02d}/{int(ano)} para conferência e autorização. Você pode baixar o Excel organizado na página:\n\n{link}"
    return link, "https://wa.me/?text=" + quote(mensagem, safe="")


def _texto(cell, valor):
    cell.value = str(valor or "")
    # Nomes e observações são texto, mesmo quando começam com =, +, - ou @.
    cell.data_type = "s"


def _cabecalho(ws, titulo, competencia, status, ultima):
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ultima)
    ws.cell(1, 1, titulo).font = Font(name="Calibri", size=17, bold=True, color="FFFFFF")
    for row in ws.iter_rows(min_row=1, max_row=1, max_col=ultima):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor="17365D")
    ws.row_dimensions[1].height = 32
    ws.cell(2, 1, f"Competência: {competencia} | {status}").font = Font(size=11, bold=True, color="17365D")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ultima)
    ws.cell(3, 1, "Dados do retrato enviado ao gestor; alterações posteriores no cadastro não modificam este relatório.").font = Font(size=10, italic=True, color="666666")
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=ultima)
    ws.freeze_panes = "B6"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:5"


def _tabela(ws, headers, dados, nome):
    for col, h in enumerate(headers, 1):
        ws.cell(5, col, h).font = Font(bold=True, color="FFFFFF")
        ws.cell(5, col).fill = PatternFill("solid", fgColor="245A81")
        ws.cell(5, col).alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[5].height = 32
    for row_idx, linha in enumerate(dados, 6):
        for col_idx, v in enumerate(linha, 1):
            cell = ws.cell(row_idx, col_idx)
            if isinstance(v, str):
                _texto(cell, v)
            else:
                cell.value = v
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[row_idx].height = 34
    if dados:
        tabela = Table(displayName=nome, ref=f"A5:{get_column_letter(len(headers))}{5 + len(dados)}")
        tabela.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tabela)
    else:
        ws.cell(6, 1, "Nenhum colaborador neste grupo.")


def excel_gestor(reg):
    if reg.get("status") not in STATUS:
        raise ValueError("Apuração indisponível para exportação: envio cancelado ou ainda não realizado.")
    snap = reg.get("snapshot") or {}
    linhas = list(snap.get("linhas") or [])
    competencia = f"{int(reg['mes']):02d}/{int(reg['ano'])}"
    status = STATUS[reg["status"]]
    grupos = {"Revisão": [], "Demais frentes": [], "Sem pagamento": []}
    for linha in linhas:
        valor = valor_numero(linha.get("GRATIFICAÇÃO"))
        frente = str(linha.get("FRENTE") or "").strip().upper()
        grupo = "Sem pagamento" if valor is None or valor <= 0 else ("Revisão" if frente in ("REVISÃO", "REVISAO") else "Demais frentes")
        grupos[grupo].append(linha)
    wb = Workbook()
    resumo = wb.active
    resumo.title = "Resumo"
    _cabecalho(resumo, "10 Sul | Apuração de gratificação", competencia, status, 5)
    dados_resumo = [[grupo, len(regs), sum((valor_numero(r.get("GRATIFICAÇÃO")) or 0) for r in regs)] for grupo, regs in grupos.items()]
    dados_resumo.append(["TOTAL", len(linhas), sum(r[2] for r in dados_resumo)])
    _tabela(resumo, ["GRUPO", "COLABORADORES", "GRATIFICAÇÃO (R$)"], dados_resumo, "ResumoGrupos")
    empresas = defaultdict(lambda: [0, 0.0])
    for linha in linhas:
        empresa = str(linha.get("EMPRESA") or "Não informada")
        empresas[empresa][0] += 1
        empresas[empresa][1] += valor_numero(linha.get("GRATIFICAÇÃO")) or 0
    resumo.cell(12, 1, "TOTAIS POR EMPRESA").font = Font(bold=True, color="17365D")
    for row, (empresa, (qtd, total)) in enumerate(sorted(empresas.items()), 13):
        _texto(resumo.cell(row, 1), empresa)
        resumo.cell(row, 2, qtd)
        resumo.cell(row, 3, total).number_format = MOEDA
    for row in range(6, 10):
        resumo.cell(row, 3).number_format = MOEDA
    for cell in resumo[9][:3]:
        cell.font = Font(bold=True, color="17365D")
        cell.fill = PatternFill("solid", fgColor="E7EFF7")
    linha_meta = 15 + len(empresas)
    _texto(resumo.cell(linha_meta, 1), "Enviado por: " + str(reg.get("enviado_por") or snap.get("enviado_por") or ""))
    _texto(resumo.cell(linha_meta + 1, 1), "Enviado em: " + str(reg.get("enviado_em") or snap.get("enviado_em") or ""))
    if reg.get("autorizado_em"):
        _texto(resumo.cell(linha_meta + 2, 1), "Autorizado em: " + str(reg["autorizado_em"]))
    for col, largura in {"A": 40, "B": 22, "C": 25, "D": 25, "E": 25}.items():
        resumo.column_dimensions[col].width = largura
    for idx, (grupo, regs) in enumerate(grupos.items(), 1):
        ws = wb.create_sheet(grupo)
        _cabecalho(ws, "10 Sul | " + grupo, competencia, status, len(COLUNAS))
        registros = []
        for r in sorted(regs, key=lambda r: (str(r.get("EMPRESA") or ""), str(r.get("COLABORADOR") or ""))):
            registros.append([valor_numero(r.get(c)) if c in ("SALÁRIO", "MÉDIA TEMPO (h)", "FALTAS", "ATESTADOS (dias)", "DNA", "DESVIOS", "INTEGRAL", "GRATIFICAÇÃO") else str(r.get(c) or "") for c in COLUNAS])
        _tabela(ws, COLUNAS, registros, f"ApuracaoGrupo{idx}")
        for col, nome in enumerate(COLUNAS, 1):
            ws.column_dimensions[get_column_letter(col)].width = 46 if nome == "COLABORADOR" else 65 if nome == "MOTIVO / CÁLCULO" else 28 if nome in ("EMPRESA", "FUNÇÃO", "STATUS") else 19
            if nome in ("SALÁRIO", "INTEGRAL", "GRATIFICAÇÃO", "MÉDIA TEMPO (h)", "ATESTADOS (dias)"):
                for row in range(6, 6 + len(registros)):
                    ws.cell(row, col).number_format = MOEDA if nome in ("SALÁRIO", "INTEGRAL", "GRATIFICAÇÃO") else "0.00"
        for row in range(6, 6 + len(registros)):
            linhas_texto = max(
                max(1, sum(max(1, ceil(len(parte) / max(1, ws.column_dimensions[cell.column_letter].width - 3))) for parte in str(cell.value or "").split("\n")))
                for cell in ws[row]
            )
            ws.row_dimensions[row].height = max(30, 15 * linhas_texto + 8)
        total_row = 7 + len(registros)
        ws.cell(total_row, 1, "TOTAL DO GRUPO").font = Font(bold=True, color="17365D")
        ws.cell(total_row, 13, sum(valor_numero(r.get("GRATIFICAÇÃO")) or 0 for r in regs)).number_format = MOEDA
        ws.cell(total_row, 13).font = Font(bold=True, color="17365D")
    saida = BytesIO()
    wb.save(saida)
    return saida.getvalue()
