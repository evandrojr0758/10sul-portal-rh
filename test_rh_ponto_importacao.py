import io
import unittest
from datetime import date, time, timedelta
from pathlib import Path
from urllib.parse import parse_qs
import pandas as pd
from openpyxl import load_workbook
import rh_ponto_importacao as m

SOURCES = Path(__file__).resolve().parent.parent / 'ponto_sources'

class TestPonto(unittest.TestCase):
 def test_durations(self):
  for v,n in [(None,0),('',0),('02:09',129),(time(1,30),90),(timedelta(hours=13),780),(.5,720),('24:00',1440)]: self.assertEqual(m.minutos(v),n)
  for v in ['FALTA','1.5','12:99',-1,'25:00']:
   with self.assertRaises(ValueError):m.minutos(v)
  self.assertEqual(m.hhmm(48*60+15),'48:15')
 @unittest.skipUnless(SOURCES.exists(), 'Planilhas reais disponíveis somente na validação local')
 def test_real_sources(self):
  total=0
  for p in SOURCES.glob('*.xlsx'):
   pack=m.ler_arquivo(p.name,p.read_bytes(),'Aracruz'); raw=pd.read_excel(p)
   raw=raw[raw['Dt. Ponto'].notna() & raw.Nome.notna()]
   raw=raw[~raw.Nome.astype(str).map(m.normalizar).isin(['TOTAL','TOTAIS','TOTAL GERAL'])]
   def minute(v):
    if pd.isna(v):return 0
    if isinstance(v,time):return v.hour*60+v.minute+v.second/60
    if isinstance(v,timedelta):return v.total_seconds()/60
    return float(v)*1440
   cols=[c for c in raw.columns if c.startswith('H. Extra') or c=='E. Feriado']
   expected=round(sum(raw[c].map(minute).sum() for c in cols));actual=sum(r['minutos_total'] for r in pack['registros'])
   self.assertEqual(actual,expected,p.name);self.assertEqual(len(pack['registros']),len(raw));total+=actual
  self.assertEqual(total,509*60+8)
  p=next(SOURCES.glob('*.csv'));pack=m.ler_arquivo(p.name,p.read_bytes(),'Aracruz')
  self.assertEqual(len(pack['registros']),62)
  self.assertTrue(all(r['minutos_total'] is None and not r['extras_informadas'] for r in pack['registros']));self.assertTrue(pack['avisos'])
 def test_csv_keys_and_errors(self):
  raw='Empresa;Dt. Ponto;Crachá;Nome;H. Extra;E. Feriado\n10 SUL SERVICE LTDA;01/10/2026;0001;JOÃO;02:00;01:00\n'
  for enc in ['cp1252','utf-8-sig']:
   p=m.ler_arquivo('a.csv',raw.encode(enc),'Aracruz');self.assertEqual(p['registros'][0]['minutos_total'],180);self.assertEqual(p['registros'][0]['cracha'],'1')
  self.assertEqual(len(m.ler_arquivo('a.csv',(raw+raw.splitlines()[1]+'\n').encode(),'Aracruz')['registros']),1)
  for changed in [raw.replace('01/10/2026','bad'),raw.replace('02:00','bad'),raw+raw.splitlines()[1].replace('02:00','03:00')+'\n']:
   with self.assertRaises(ValueError):m.ler_arquivo('a.csv',changed.encode(),'Aracruz')
  a=m.ler_arquivo('a.csv',raw.encode(),'Aracruz')['registros'][0]
  b=m.ler_arquivo('b.csv',raw.replace('0001','1').replace('02:00','03:00').encode(),'Aracruz')['registros'][0]
  self.assertEqual(a['chave'],b['chave'])
  self.assertNotEqual(a['chave'],m.ler_arquivo('a.csv',raw.encode(),'Bahia')['registros'][0]['chave'])
 def test_pagination(self):
  calls=[]
  def sb(method,table,params):
   p=parse_qs(params);calls.append(p);self.assertEqual(p['data_ponto'],['gte.2026-07-01','lte.2026-10-28'])
   return [{'chave':str(i)} for i in range(1000)] if p['offset']==['0'] else [{'chave':'last'}]
  self.assertEqual(len(m.carregar_base(sb,date(2026,7,1),date(2026,10,28),'Aracruz')),1001);self.assertEqual(len(calls),2)
 def test_summaries_and_excel(self):
  rows=[]
  for dia,valor,disponivel in [('2026-10-01',900,True),('2026-10-02',900,True),('2026-10-03',None,False)]:
   row=dict(data_ponto=dia,empresa='EMPRESA TESTE',cracha='1',colaborador='COLABORADOR TESTE',extras_informadas=disponivel,minutos_total=valor,arquivo='ponto.xlsx')
   row.update({c:0 for c in m.COMPONENTES});row['minutos_extra']=valor or 0;rows.append(row)
  df=pd.DataFrame(rows);grupos=m.resumos(df)
  for g in grupos.values():self.assertEqual(g.minutos_total.sum(),1800)
  wb=load_workbook(io.BytesIO(m.excel_resumo(df,grupos)));self.assertEqual(wb.sheetnames,['Mensal','Diario','Colaborador','Detalhamento'])
  ws=wb['Colaborador'];heads=[c.value for c in ws[1]];i=heads.index('Total de extras')+1;self.assertEqual(ws.cell(2,i).number_format,'[h]:mm')
  ws=wb['Detalhamento'];heads=[c.value for c in ws[1]];flag=heads.index('Extras informadas')+1;tot=heads.index('Total de extras')+1
  self.assertTrue(all(ws.cell(r,tot).value is None for r in range(2,ws.max_row+1) if ws.cell(r,flag).value is False))

UI_HISTORICO = '''
import streamlit as st
import rh_ponto_importacao as m
st.session_state['usuario_logado']={'perfil':'RH','nome':'Teste'}
if 'historico_teste' not in st.session_state:
    st.session_state['historico_teste']=[dict(id='a',arquivo='planilha_a.xlsx',unidade='Aracruz',registros=10,importado_por='Teste',criado_em='2026-10-29',excluido_em=None,excluido_por=None),dict(id='b',arquivo='planilha_b.xlsx',unidade='Aracruz',registros=20,importado_por='Teste',criado_em='2026-10-28',excluido_em=None,excluido_por=None)]
def sb(method,table,params='',payload=None):
    if method=='GET':return st.session_state['historico_teste']
    assert table=='rpc/rh_excluir_importacao_ponto'
    for r in st.session_state['historico_teste']:
        if r['id']==payload['p_id']:
            r['excluido_em']='2026-10-29';r['excluido_por']=payload['p_usuario']
            return {'excluido':True}
    raise ValueError('Não encontrado')
m.tela_historico(sb,'Aracruz')
'''

class TestExcluirPonto(unittest.TestCase):
 def test_payload_exclusao(self):
  calls=[]
  m.excluir_importacao(lambda *a:calls.append(a),'a','Aracruz','RH Teste')
  self.assertEqual(calls[0],('POST','rpc/rh_excluir_importacao_ponto','',{'p_id':'a','p_unidade':'Aracruz','p_usuario':'RH Teste'}))
 def test_confirmacao_e_resumo(self):
  from streamlit.testing.v1 import AppTest
  at=AppTest.from_string(UI_HISTORICO,default_timeout=20).run()
  self.assertFalse(at.exception);self.assertEqual(len(at.button),0)
  at.selectbox[0].set_value('a').run()
  self.assertTrue(at.button[0].disabled)
  self.assertIn('planilha_a.xlsx',at.warning[0].value)
  at.checkbox(key='ponto_excluir_confirmacao_a').check().run()
  self.assertFalse(at.button[0].disabled)
  at.button[0].click().run()
  self.assertFalse(at.exception)
  self.assertIsNotNone(at.session_state['historico_teste'][0]['excluido_em'])
  self.assertIsNone(at.session_state['historico_teste'][1]['excluido_em'])
  self.assertTrue(any('Resumos atualizados' in msg.value for msg in at.success))
 def test_troca_de_arquivo_exige_nova_confirmacao(self):
  from streamlit.testing.v1 import AppTest
  at=AppTest.from_string(UI_HISTORICO,default_timeout=20).run()
  at.selectbox[0].set_value('a').run()
  at.checkbox(key='ponto_excluir_confirmacao_a').check().run()
  at.selectbox[0].set_value('b').run()
  self.assertTrue(at.button[0].disabled)
  self.assertFalse(at.checkbox(key='ponto_excluir_confirmacao_b').value)
  self.assertTrue(all(r['excluido_em'] is None for r in at.session_state['historico_teste']))

if __name__=='__main__':unittest.main()
