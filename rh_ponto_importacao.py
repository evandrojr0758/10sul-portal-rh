"""Importação persistente do ponto eletrônico e resumos das horas extras informadas."""
import csv
import hashlib
import io
import json
import math
import re
import unicodedata
from datetime import date, datetime, time, timedelta

import pandas as pd
import streamlit as st
from openpyxl.styles import Font, PatternFill

COMPONENTES = {
    'minutos_extra': ('H EXTRA', 'H EXTRAS', 'HORA EXTRA', 'HORAS EXTRAS'),
    'minutos_extra_noturna': ('H EXTRA NOT', 'H EXTRA NOTURNA'),
    'minutos_extra_folga': ('H EXTRA FOLGA',),
    'minutos_extra_folga_noturna': ('H EXTRA F NOT', 'H EXTRA FOLGA NOT'),
    'minutos_extra_feriado': ('E FERIADO', 'H EXTRA FERIADO'),
}
ROTULOS = {'minutos_extra': 'Extra', 'minutos_extra_noturna': 'Extra noturna',
           'minutos_extra_folga': 'Extra folga', 'minutos_extra_folga_noturna': 'Extra folga noturna',
           'minutos_extra_feriado': 'Extra feriado', 'minutos_total': 'Total de extras'}


def normalizar(valor):
    return ' '.join(re.sub(r'[^A-Z0-9]+', ' ', ''.join(
        c for c in unicodedata.normalize('NFD', str(valor or '').upper())
        if not unicodedata.combining(c))).split())


def vazio(valor):
    return valor is None or (not isinstance(valor, (list, dict)) and bool(pd.isna(valor))) or str(valor).strip() == ''


def minutos(valor):
    """Duração Excel ou HH:MM[:SS]; nunca interpreta horário decimal como HH.MM."""
    if vazio(valor) or str(valor).strip() in ('-', '--'):
        return 0
    if isinstance(valor, (timedelta, pd.Timedelta)):
        total = valor.total_seconds() / 60
    elif isinstance(valor, time):
        total = valor.hour * 60 + valor.minute + valor.second / 60
    elif isinstance(valor, (float, int)) and not isinstance(valor, bool):
        total = float(valor) * 1440
    else:
        match = re.fullmatch(r'(\d+):(\d{2})(?::(\d{2}))?', str(valor).strip())
        if not match or int(match[2]) >= 60 or int(match[3] or 0) >= 60:
            raise ValueError(f'Duração inválida: {valor!s}. Use HH:MM.')
        total = int(match[1]) * 60 + int(match[2]) + int(match[3] or 0) / 60
    if not math.isfinite(total) or total < 0 or total > 1440:
        raise ValueError(f'Duração diária fora do intervalo de 00:00 a 24:00: {valor!s}')
    return int(round(total))


def hhmm(valor):
    if vazio(valor):
        return 'Não informado'
    n = int(round(float(valor)))
    return f'{n // 60:02d}:{n % 60:02d}'


def _texto(valor):
    return '' if vazio(valor) else str(valor).strip()


def _cracha(valor):
    texto = _texto(valor)
    if re.fullmatch(r'\d+(?:\.0+)?', texto):
        return str(int(float(texto))) if '.' in texto else (texto.lstrip('0') or '0')
    return texto.upper()


def _empresa(valor):
    texto = normalizar(valor)
    if texto.startswith('10 SUL'):
        if 'PRESTADORA' in texto:
            return '10 SUL PRESTADORA'
        if 'SERVICE' in texto:
            return '10 SUL SERVICE'
    return texto


def _json_cell(valor):
    if vazio(valor):
        return None
    if isinstance(valor, (datetime, date, time)):
        return valor.isoformat()
    if isinstance(valor, timedelta):
        return str(valor)
    if hasattr(valor, 'item'):
        return valor.item()
    return valor if isinstance(valor, (str, int, float, bool)) else str(valor)


def _abas(nome, conteudo):
    if nome.lower().endswith('.csv'):
        for encoding in ('utf-8-sig', 'cp1252'):
            try:
                texto = conteudo.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError('Não foi possível ler a codificação do CSV.')
        try:
            sep = csv.Sniffer().sniff(texto[:12000], delimiters=';,\t').delimiter
        except csv.Error:
            sep = ';'
        return {'CSV': pd.read_csv(io.StringIO(texto), sep=sep, header=None, dtype=object, keep_default_na=False)}
    if nome.lower().endswith(('.xlsx', '.xls')):
        return pd.read_excel(io.BytesIO(conteudo), sheet_name=None, header=None,
                             engine='xlrd' if nome.lower().endswith('.xls') else 'openpyxl')
    raise ValueError('Envie arquivos XLSX, XLS ou CSV.')


def ler_arquivo(nome, conteudo, unidade):
    registros, avisos, vistas = [], [], {}
    for aba, bruto in _abas(nome, conteudo).items():
        cab = None
        for i in range(min(len(bruto), 40)):
            nomes = [normalizar(_texto(x)) for x in bruto.iloc[i]]
            if ('NOME' in nomes or 'COLABORADOR' in nomes) and ('DT PONTO' in nomes or 'DATA' in nomes):
                cab = i
                break
        if cab is None:
            avisos.append(f'Aba {aba}: sem tabela de ponto reconhecida.')
            continue
        headers = [_texto(x) for x in bruto.iloc[cab]]
        mapa = {normalizar(x): i for i, x in enumerate(headers) if x}
        col_nome = mapa.get('NOME', mapa.get('COLABORADOR'))
        col_data = mapa.get('DT PONTO', mapa.get('DATA'))
        col_empresa, col_cracha = mapa.get('EMPRESA'), mapa.get('CRACHA')
        cols_extra = {campo: next((mapa[a] for a in aliases if a in mapa), None) for campo, aliases in COMPONENTES.items()}
        disponivel = any(i is not None for i in cols_extra.values())
        if not disponivel:
            avisos.append(f'Aba {aba}: sem colunas de horas extras. Os registros serão guardados como não informados e não entrarão no total de extras.')
        for idx in range(cab + 1, len(bruto)):
            valores = bruto.iloc[idx].tolist()
            nome_col = _texto(valores[col_nome])
            data_raw = valores[col_data]
            if not nome_col and vazio(data_raw):
                continue
            if normalizar(nome_col) in ('TOTAL', 'TOTAIS', 'TOTAL GERAL'):
                continue
            if not nome_col:
                raise ValueError(f'{aba}, linha {idx + 1}: colaborador vazio.')
            try:
                if vazio(data_raw) or isinstance(data_raw, (int, float)):
                    raise ValueError('data vazia ou não reconhecida')
                dia = pd.to_datetime(data_raw, dayfirst=isinstance(data_raw, str) and '/' in data_raw, errors='raise').date()
                if dia.year < 1900 or dia.year > 2100:
                    raise ValueError('ano inválido')
            except (ValueError, TypeError, OverflowError) as exc:
                raise ValueError(f'{aba}, linha {idx + 1}: data do ponto inválida ({data_raw!s}).') from exc
            empresa = _empresa(valores[col_empresa]) if col_empresa is not None else ''
            if not empresa:
                raise ValueError(f'{aba}, linha {idx + 1}: informe a empresa na planilha.')
            cracha = _cracha(valores[col_cracha]) if col_cracha is not None else ''
            identidade = 'CRACHA:' + cracha if cracha else 'NOME:' + normalizar(nome_col)
            chave = hashlib.sha256(json.dumps([unidade, empresa, identidade, dia.isoformat()], ensure_ascii=False).encode()).hexdigest()
            r = dict(chave=chave, data_ponto=dia.isoformat(), empresa=empresa, cracha=cracha,
                     colaborador=nome_col.upper(), extras_informadas=disponivel,
                     aba=aba, linha=idx + 1,
                     dados_originais={h: _json_cell(v) for h, v in zip(headers, valores) if h})
            try:
                r.update({campo: minutos(valores[i]) if i is not None else 0 for campo, i in cols_extra.items()})
            except ValueError as exc:
                raise ValueError(f'{aba}, linha {idx + 1}, {nome_col}: {exc}') from exc
            r['minutos_total'] = sum(r[c] for c in COMPONENTES) if disponivel else None
            # Duplicatas idênticas são ignoradas; divergências exigem conferência.
            assinatura = {k: v for k, v in r.items() if k not in ('aba', 'linha')}
            if chave in vistas:
                if vistas[chave] != assinatura:
                    raise ValueError(f'{aba}, linha {idx + 1}: registros divergentes para {nome_col} em {dia:%d/%m/%Y}.')
                continue
            vistas[chave] = assinatura
            registros.append(r)
    if not registros:
        raise ValueError('Nenhum registro de ponto encontrado. Verifique as colunas Empresa, Nome e Dt. Ponto.')
    if len(registros) > 50000:
        raise ValueError('Limite de 50 mil registros por arquivo. Divida a exportação por período.')
    return {'arquivo': nome, 'arquivo_hash': hashlib.sha256(conteudo).hexdigest(),
            'registros': registros, 'avisos': avisos}


def salvar_arquivo(sb, pacote, unidade, usuario):
    return sb('POST', 'rpc/rh_importar_ponto', '', {
        'p_arquivo': pacote['arquivo'], 'p_hash': pacote['arquivo_hash'],
        'p_unidade': unidade, 'p_usuario': str(usuario), 'p_registros': pacote['registros']})


def carregar_importacoes(sb, unidade):
    from urllib.parse import urlencode
    historico, offset = [], 0
    while True:
        params = urlencode({'select': 'id,arquivo,unidade,registros,importado_por,criado_em,excluido_em,excluido_por',
                            'unidade': 'eq.' + unidade, 'order': 'criado_em.desc,id.desc',
                            'limit': 1000, 'offset': offset})
        lote = sb('GET', 'rh_ponto_importacoes', params) or []
        historico.extend(lote)
        if len(lote) < 1000:
            return historico
        offset += len(lote)


def excluir_importacao(sb, importacao_id, unidade, usuario):
    return sb('POST', 'rpc/rh_excluir_importacao_ponto', '', {
        'p_id': importacao_id, 'p_unidade': unidade, 'p_usuario': str(usuario)})


def tela_historico(sb, unidade):
    mensagem = st.session_state.pop('ponto_exclusao_sucesso', None)
    if mensagem:
        st.success(mensagem)
    try:
        historico = carregar_importacoes(sb, unidade)
    except Exception as exc:
        st.error(f'Não foi possível carregar o histórico: {exc}'); return
    if not historico:
        st.info('Nenhuma planilha importada nesta unidade.'); return
    mostrar_excluidas = st.checkbox('Mostrar planilhas excluídas', key='ponto_mostrar_excluidas_' + unidade)
    visiveis = [r for r in historico if mostrar_excluidas or not r.get('excluido_em')]
    tabela = pd.DataFrame(visiveis)
    if not tabela.empty:
        tabela['Situação'] = ['Excluída' if r.get('excluido_em') else 'Ativa' for r in visiveis]
        st.dataframe(tabela.drop(columns=['id'], errors='ignore').rename(columns={
            'arquivo': 'Arquivo', 'unidade': 'Unidade', 'registros': 'Registros',
            'importado_por': 'Importado por', 'criado_em': 'Importado em',
            'excluido_em': 'Excluído em', 'excluido_por': 'Excluído por'}), hide_index=True, use_container_width=True)
    ativos = {r['id']: r for r in historico if not r.get('excluido_em')}
    if not ativos:
        st.info('Não há planilhas ativas para excluir.'); return
    selecionado = st.selectbox('Planilha para excluir', list(ativos), index=None,
        placeholder='Selecione uma importação...', key='ponto_excluir_selecao_' + unidade,
        format_func=lambda rid: f"{ativos[rid]['arquivo']} — {ativos[rid]['registros']} registros — {ativos[rid]['criado_em']}")
    if selecionado is None:
        return
    registro = ativos[selecionado]
    st.warning(f"Excluir a planilha {registro['arquivo']} ({registro['registros']} registros) de {unidade}?")
    st.caption('Os registros desta importação sairão dos resumos. Se outra planilha contiver os mesmos dias, serão usados os dados dela. Você poderá importar este arquivo novamente depois.')
    confirmado = st.checkbox('Confirmo a exclusão desta planilha da apuração', key='ponto_excluir_confirmacao_' + selecionado)
    if st.button('🗑️ Excluir planilha importada', type='primary', disabled=not confirmado, key='ponto_excluir_botao_' + selecionado):
        try:
            usr = st.session_state.get('usuario_logado') or {}
            if str(usr.get('perfil') or '').upper() not in ('RH', 'ADMIN'):
                raise ValueError('Acesso permitido somente ao RH e Administrador.')
            excluir_importacao(sb, selecionado, unidade, usr.get('nome') or usr.get('usuario') or usr.get('id'))
        except Exception as exc:
            st.error(f'Não foi possível excluir: {exc}')
        else:
            st.session_state['ponto_exclusao_sucesso'] = f"Planilha {registro['arquivo']} excluída. Resumos atualizados."
            st.rerun()


def carregar_base(sb, inicio, fim, unidade):
    from urllib.parse import urlencode
    rows, offset = [], 0
    while True:
        params = urlencode([('select', '*'), ('data_ponto', f'gte.{inicio.isoformat()}'),
                            ('data_ponto', f'lte.{fim.isoformat()}'), ('unidade', f'eq.{unidade}'),
                            ('order', 'data_ponto.asc,chave.asc'), ('limit', 1000), ('offset', offset)])
        lote = sb('GET', 'rh_ponto_atual', params) or []
        rows.extend(lote)
        if len(lote) < 1000:
            return pd.DataFrame(rows)
        offset += len(lote)


def resumos(df):
    base = df[df['extras_informadas'].eq(True)].copy()
    base['MES'] = pd.to_datetime(base['data_ponto']).dt.strftime('%Y-%m')
    campos = list(COMPONENTES) + ['minutos_total']
    return {
        'Mensal': base.groupby('MES', as_index=False)[campos].sum(),
        'Diario': base.groupby('data_ponto', as_index=False)[campos].sum(),
        'Colaborador': base.groupby(['empresa', 'cracha', 'colaborador'], as_index=False)[campos].sum().sort_values('minutos_total', ascending=False),
    }


def exibir_horas(df):
    out = df.copy()
    for campo, rotulo in ROTULOS.items():
        if campo in out:
            out[campo] = out[campo].map(hhmm)
    return out.rename(columns={**ROTULOS, 'data_ponto': 'Data', 'MES': 'Mês', 'empresa': 'Empresa',
                              'cracha': 'Crachá', 'colaborador': 'Colaborador', 'arquivo': 'Arquivo',
                              'extras_informadas': 'Extras informadas'})


def excel_resumo(df, grupos):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        abas = {**grupos, 'Detalhamento': df[[c for c in ['data_ponto', 'empresa', 'cracha', 'colaborador', 'extras_informadas', *ROTULOS, 'arquivo'] if c in df]].copy()}
        for nome, dados in abas.items():
            dados = dados.copy()
            for campo in ROTULOS:
                if campo in dados:
                    dados[campo] = pd.to_numeric(dados[campo], errors='coerce') / 1440
            if 'data_ponto' in dados:
                dados['data_ponto'] = pd.to_datetime(dados['data_ponto']).dt.date
            dados = dados.rename(columns={**ROTULOS, 'data_ponto': 'Data', 'MES': 'Mês', 'empresa': 'Empresa', 'cracha': 'Crachá', 'colaborador': 'Colaborador', 'arquivo': 'Arquivo', 'extras_informadas': 'Extras informadas'})
            dados.to_excel(writer, sheet_name=nome, index=False)
            ws = writer.sheets[nome]
            ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='243B53')
            for col in ws.columns:
                label = col[0].value
                ws.column_dimensions[col[0].column_letter].width = 43 if label == 'Colaborador' else (50 if label == 'Arquivo' else max(17, len(str(label)) + 2))
                for cell in col[1:]:
                    if label in ROTULOS.values(): cell.number_format = '[h]:mm'
                    if label == 'Data': cell.number_format = 'DD/MM/YYYY'
                    if cell.data_type == 'f': cell.data_type = 's'
    return output.getvalue()


def tela(sb):
    if str((st.session_state.get('usuario_logado') or {}).get('perfil') or '').upper() not in ('RH', 'ADMIN'):
        st.error('Acesso permitido somente ao RH e Administrador.'); return
    st.title('⏱️ Apuração de horas extras')
    st.caption('Importe o ponto eletrônico. O histórico fica salvo para consultas por mês, dia e colaborador.')
    unidade = st.selectbox('Unidade', ['Aracruz', 'Bahia', 'Três Lagoas'], key='ponto_unidade')
    with st.expander('📤 Importar planilhas de ponto', expanded=True):
        arquivos = st.file_uploader('Selecione uma ou mais planilhas', type=['xlsx', 'xls', 'csv'], accept_multiple_files=True, key='ponto_upload')
        st.caption('São usadas as horas extras informadas pelo ponto: Extra, Extra noturna, Extra folga, Extra folga noturna e Extra feriado. O mesmo arquivo não duplica registros; uma nova exportação atualiza os dias correspondentes, mantendo os anteriores no histórico.')
        pacotes, erros = [], []
        for arquivo in arquivos or []:
            try:
                pacote = ler_arquivo(arquivo.name, arquivo.getvalue(), unidade)
                pacotes.append(pacote)
                regs = pacote['registros']
                st.write(f"**{arquivo.name}** — {len(regs)} registro(s), de {min(r['data_ponto'] for r in regs)} a {max(r['data_ponto'] for r in regs)}")
                for aviso in pacote['avisos']: st.warning(aviso)
            except Exception as exc:
                erros.append(arquivo.name); st.error(f'{arquivo.name}: {exc}')
        if st.button('💾 Importar e salvar na base', type='primary', disabled=not pacotes or bool(erros), key='ponto_salvar'):
            for pacote in pacotes:
                try:
                    usr = st.session_state.get('usuario_logado') or {}
                    resultado = salvar_arquivo(sb, pacote, unidade, usr.get('nome') or usr.get('usuario') or usr.get('id'))
                    if resultado.get('duplicado'):
                        st.info(f"{pacote['arquivo']}: já importado nesta unidade; nenhum registro duplicado.")
                    else:
                        st.success(f"{pacote['arquivo']}: {resultado['registros']} registro(s) salvos.")
                except Exception as exc:
                    st.error(f"Não foi possível salvar {pacote['arquivo']}: {exc}"); break
    st.subheader('Resumo das horas extras')
    a, b = st.columns(2)
    hoje = date.today()
    inicio = a.date_input('De', date(hoje.year, hoje.month, 1), key='ponto_inicio', format='DD/MM/YYYY')
    fim = b.date_input('Até', hoje, key='ponto_fim', format='DD/MM/YYYY')
    if inicio > fim:
        st.warning('A data inicial deve ser anterior ou igual à final.'); return
    try:
        df = carregar_base(sb, inicio, fim, unidade)
    except Exception as exc:
        st.error(f'Não foi possível consultar a base: {exc}'); return
    with st.expander('📋 Histórico das importações / Excluir planilha'):
        tela_historico(sb, unidade)
    if df.empty:
        st.info('Nenhum ponto importado para esta unidade e período. Importe as planilhas acima ou altere as datas.'); return
    empresa = st.selectbox('Empresa', ['Todas'] + sorted(df['empresa'].dropna().unique().tolist()), key='ponto_empresa')
    if empresa != 'Todas': df = df[df['empresa'].eq(empresa)]
    pessoas = sorted(df['colaborador'].dropna().unique().tolist())
    selecionados = st.multiselect('Colaboradores (vazio = todos)', pessoas, key='ponto_pessoas')
    if selecionados: df = df[df['colaborador'].isin(selecionados)]
    if df.empty:
        st.info('Nenhum registro com os filtros selecionados.'); return
    desconhecidos = int((~df['extras_informadas']).sum())
    if desconhecidos:
        st.warning(f'{desconhecidos} registro(s) sem colunas de horas extras. Não entram nos totais; importe a exportação completa do ponto para apurá-los.')
    grupos = resumos(df)
    conhecidos = df[df['extras_informadas']]
    m1, m2, m3 = st.columns(3)
    m1.metric('Total de horas extras informadas', hhmm(conhecidos['minutos_total'].sum()) if not conhecidos.empty else 'Não informado')
    m2.metric('Colaboradores com horas extras', int(conhecidos.loc[conhecidos['minutos_total'].gt(0), 'colaborador'].nunique()))
    m3.metric('Dias com horas extras', int(conhecidos.loc[conhecidos['minutos_total'].gt(0), 'data_ponto'].nunique()))
    for titulo, chave, coluna in [('Horas extras por mês', 'Mensal', 'MES'), ('Horas extras por dia', 'Diario', 'data_ponto'), ('Horas extras por colaborador', 'Colaborador', 'colaborador')]:
        st.subheader(titulo)
        resumo = grupos[chave]
        if resumo.empty:
            st.info('Sem horas extras informadas neste recorte.'); continue
        grafico = resumo.groupby(coluna)['minutos_total'].sum().div(60).rename('Horas extras (decimal)')
        st.bar_chart(grafico, use_container_width=True)
        st.dataframe(exibir_horas(resumo), hide_index=True, use_container_width=True)
    with st.expander('Detalhamento dos registros'):
        cols = ['data_ponto', 'empresa', 'cracha', 'colaborador', 'extras_informadas', *ROTULOS, 'arquivo']
        st.dataframe(exibir_horas(df[[c for c in cols if c in df]]), hide_index=True, use_container_width=True)
    st.download_button('📥 Baixar resumo em Excel', excel_resumo(df, grupos),
                       file_name=f'Horas_Extras_{unidade.replace(" ", "_")}_{inicio}_{fim}.xlsx',
                       mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', key='ponto_excel')
