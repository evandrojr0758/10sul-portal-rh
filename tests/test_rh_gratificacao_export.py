import unittest
from io import BytesIO
from urllib.parse import parse_qs, urlparse
from openpyxl import load_workbook
from rh_gratificacao_export import excel_gestor, whatsapp_aprovacao


class ExportacaoGestorTest(unittest.TestCase):
    def registro(self):
        return {"ano": 2026, "mes": 10, "status": "AGUARDANDO_AUTORIZACAO", "snapshot": {"linhas": [
            {"COLABORADOR": "REVISÃO TESTE", "EMPRESA": "SERVICE", "FRENTE": "REVISÃO", "GRATIFICAÇÃO": 100, "SALÁRIO": 3000},
            {"COLABORADOR": "ITR TESTE", "EMPRESA": "PRESTADORA", "FRENTE": "ITR", "GRATIFICAÇÃO": 360},
            {"COLABORADOR": "=FORMULA()", "EMPRESA": "SERVICE", "FRENTE": "FABRICAÇÃO", "GRATIFICAÇÃO": 0},
            {"COLABORADOR": "PENDENTE", "EMPRESA": "SERVICE", "FRENTE": "ITR", "GRATIFICAÇÃO": None},
        ]}}

    def test_abas_totais_formatacao(self):
        wb = load_workbook(BytesIO(excel_gestor(self.registro())))
        self.assertEqual(wb.sheetnames, ["Resumo", "Revisão", "Demais frentes", "Sem pagamento"])
        self.assertEqual(wb["Resumo"]["B9"].value, 4)
        self.assertEqual(wb["Resumo"]["C9"].value, 460)
        self.assertEqual(wb["Revisão"]["M6"].value, 100)
        self.assertEqual(wb["Demais frentes"]["M6"].value, 360)
        self.assertEqual(wb["Sem pagamento"]["A6"].data_type, "s")
        self.assertEqual(wb["Sem pagamento"]["A6"].value, "=FORMULA()")
        self.assertEqual(wb["Revisão"].freeze_panes, "B6")
        self.assertTrue(wb["Revisão"].tables)
        self.assertIn("AGUARDANDO AUTORIZAÇÃO", wb["Resumo"]["A2"].value)

    def test_autorizada_e_cancelada(self):
        reg = self.registro()
        reg["status"] = "AUTORIZADA"
        self.assertTrue(excel_gestor(reg).startswith(b"PK"))
        reg["status"] = "CANCELADA"
        with self.assertRaises(ValueError):
            excel_gestor(reg)

    def test_sem_colaboradores(self):
        reg = self.registro()
        reg["snapshot"]["linhas"] = []
        wb = load_workbook(BytesIO(excel_gestor(reg)))
        self.assertEqual(wb["Resumo"]["C9"].value, 0)

    def test_whatsapp_mensagem_e_link(self):
        link, wa = whatsapp_aprovacao("https://exemplo.test/?outro=1", "token_123-", 2026, 10)
        self.assertEqual(link, "https://exemplo.test/?aprovar_gratificacao=token_123-")
        self.assertEqual(urlparse(wa).hostname, "wa.me")
        mensagem = parse_qs(urlparse(wa).query)["text"][0]
        self.assertIn("10/2026", mensagem)
        self.assertIn(link, mensagem)
        self.assertNotIn("phone", parse_qs(urlparse(wa).query))


if __name__ == "__main__":
    unittest.main()
